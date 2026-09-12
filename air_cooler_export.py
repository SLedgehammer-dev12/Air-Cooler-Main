import io
import json
import os
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm, cm
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer,
    PageBreak, Image
)
from reportlab.platypus.flowables import HRFlowable

from air_cooler_main_core import APP_VERSION


def _json_safe(value):
    """Dataclass, Pint Quantity ve numpy değerlerini JSON-serializable yapıya derinlemesine çevirir."""
    import dataclasses
    try:
        import numpy as np
    except ImportError:
        np = None

    if dataclasses.is_dataclass(value) and not isinstance(value, type):
        return _json_safe(dataclasses.asdict(value))
    if hasattr(value, "magnitude") and hasattr(value, "units"):
        return {"magnitude": float(value.magnitude), "unit": str(value.units)}
    if isinstance(value, dict):
        return {str(k): _json_safe(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_safe(v) for v in value]
    if np is not None:
        if isinstance(value, np.ndarray):
            return _json_safe(value.tolist())
        if isinstance(value, (np.generic, np.number)):
            return value.item()
    if isinstance(value, (int, float, str, bool)) or value is None:
        return value
    return str(value)


def _sanitize(val, default="—"):
    if val is None:
        return default
    if isinstance(val, float):
        if abs(val) < 1e-12:
            return 0.0
        return val
    return val


def _fmt(val, decimals=2):
    v = _sanitize(val)
    if isinstance(v, float):
        return f"{v:.{decimals}f}"
    return str(v)


EXPORT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "exports")
os.makedirs(EXPORT_DIR, exist_ok=True)


def _row(label, value):
    return (label, str(value))


def _build_common_sections(res, geom_params, kompozisyon, mode):
    sections = []

    proc_rows = [
        _row("Mod", mode),
        _row("Kütlesel Debisi", f"{_sanitize(res.get('m_dot_kg_s', res.get('m_dot_air_kg_s')))} kg/s"),
    ]
    if "T_in_C" in res:
        proc_rows.append(_row("Giriş Sıcaklığı", f"{_fmt(res.get('T_in_C'))} °C"))
    if "T_out_C" in res:
        proc_rows.append(_row("Çıkış Sıcaklığı", f"{_fmt(res.get('T_out_C'))} °C"))
    if "P_in_bar" in res:
        proc_rows.append(_row("Giriş Basıncı", f"{_fmt(res.get('P_in_bar'))} bar"))
    if "gas_in_phase" in res:
        proc_rows.append(_row("Giriş Fazı", _sanitize(res.get("gas_in_phase"))))
        proc_rows.append(_row("Çıkış Fazı", _sanitize(res.get("gas_out_phase"))))
    if "dP_air_Pa" in res:
        proc_rows.append(_row("Hava Basınç Kaybı", f"{_fmt(res.get('dP_air_Pa'))} Pa"))
        proc_rows.append(_row("Fan Şaft Gücü", f"{_fmt(res.get('fan_power_kW'))} kW"))
    if "gas_dP_bar" in res:
        proc_rows.append(_row("Gaz Basınç Düşümü", f"{_fmt(res.get('gas_dP_bar'))} bar"))
    sections.append(("Proses Şartları", proc_rows))

    perf_rows = []
    if "Q_kW" in res:
        perf_rows.append(_row("Toplam Isı Yükü (Q)", f"{_fmt(res.get('Q_kW'))} kW"))
    if "U_W_m2K" in res:
        perf_rows.append(_row("Toplam Isı Geçiş Katsayısı (U)", f"{_fmt(res.get('U_W_m2K'))} W/m²K"))
    if "h_inside_W_m2K" in res:
        perf_rows.append(_row("Boru İçi Film Katsayısı (hi)", f"{_fmt(res.get('h_inside_W_m2K'))} W/m²K"))
    if "h_outside_actual_W_m2K" in res:
        perf_rows.append(_row("Dış Film Katsayısı (ho)", f"{_fmt(res.get('h_outside_actual_W_m2K'))} W/m²K"))
    if "fin_efficiency" in res:
        perf_rows.append(_row("Kanat Verimi", f"{_fmt(res.get('fin_efficiency') * 100)} %"))
    if "surface_efficiency" in res:
        perf_rows.append(_row("Yüzey Verimi", f"{_fmt(res.get('surface_efficiency') * 100)} %"))
    if "actual_area_m2" in res:
        perf_rows.append(_row("Gerçek Alan", f"{_fmt(res.get('actual_area_m2'))} m²"))
    if "required_area_m2" in res:
        perf_rows.append(_row("Gerekli Alan", f"{_fmt(res.get('required_area_m2'))} m²"))
    if "overdesign_pct" in res:
        perf_rows.append(_row("Overdesign", f"{_fmt(res.get('overdesign_pct'))} %"))
    if "lmtd_K" in res:
        perf_rows.append(_row("LMTD", f"{_fmt(res.get('lmtd_K'))} K"))
    if "Ft" in res:
        perf_rows.append(_row("LMTD Düzeltme Faktörü (F)", f"{_fmt(res.get('Ft'), 4)}"))
    if "m_dot_air_kg_s" in res:
        perf_rows.append(_row("Hava Debisi", f"{_fmt(res.get('m_dot_air_kg_s'))} kg/s"))
    if "fan_tip_speed_m_s" in res:
        perf_rows.append(_row("Fan Kanat Uç Hızı", f"{_fmt(res.get('fan_tip_speed_m_s'))} m/s"))
    if "fan_sound_power_dB" in res:
        perf_rows.append(_row("Fan Ses Gücü (yaklaşık)", f"{_fmt(res.get('fan_sound_power_dB'))} dB"))
    if "condensation_applied" in res:
        perf_rows.append(_row("Yoğuşma Uygulandı", "Evet" if res.get("condensation_applied") else "Hayır"))
    if res.get("segmental_applied"):
        perf_rows.append(_row("Bölgesel (Segmental) Hesap", "Uygulandı"))
    if "gas_velocity_m_s" in res:
        perf_rows.append(_row("Boru İçi Gaz Hızı", f"{_fmt(res.get('gas_velocity_m_s'))} m/s"))
    if "gas_Re" in res:
        perf_rows.append(_row("Boru İçi Reynolds", f"{_fmt(res.get('gas_Re'), 0)}"))
    sections.append(("Performans Verileri", perf_rows))

    geom_rows = [
        _row("Sıra Sayısı", _sanitize(geom_params.get("tube_rows"))),
        _row("Geçiş Sayısı", _sanitize(geom_params.get("tube_passes"))),
        _row("Sıra Başına Tüp", _sanitize(geom_params.get("tubes_per_row"))),
        _row("Tüp Boyu", f"{_fmt(geom_params.get('tube_length'))} m"),
        _row("Tüp Dış Çapı", f"{_fmt(geom_params.get('tube_od') * 1000)} mm"),
        _row("Tüp Et Kalınlığı", f"{_fmt(geom_params.get('tube_thickness') * 1000)} mm"),
        _row("Boru Eksene Adımı", f"{_fmt(geom_params.get('pitch') * 1000)} mm"),
        _row("Dizilim Açısı", f"{_fmt(geom_params.get('angle'), 0)}°"),
        _row("Kanatçık Yüksekliği", f"{_fmt(geom_params.get('fin_height') * 1000)} mm"),
        _row("Kanatçık Kalınlığı", f"{_fmt(geom_params.get('fin_thickness') * 1000)} mm"),
        _row("Kanatçık Yoğunluğu", f"{_fmt(geom_params.get('fin_density') / 39.37)} FPI"),
    ]
    if "fin_type" in geom_params:
        geom_rows.append(_row("Kanat Bağlantı Tipi", _sanitize(geom_params.get("fin_type"))))
    sections.append(("Boru & Kanat Geometrisi", geom_rows))

    mech_rows = [
        _row("Boru Malzemesi", _sanitize(geom_params.get("tube_mat", "—"))),
        _row("Kanatçık Malzemesi", _sanitize(geom_params.get("fin_mat", "—"))),
    ]
    if "asme_grade" in geom_params:
        mech_rows.append(_row("Boru Malzeme Sınıfı (ASME)", _sanitize(geom_params.get("asme_grade"))))
    if "ca" in geom_params:
        mech_rows.append(_row("Korozyon Payı (CA)", f"{_fmt(geom_params.get('ca'))} mm"))
    if "header_type" in geom_params:
        mech_rows.append(_row("Kollektör (Header) Tipi", _sanitize(geom_params.get("header_type"))))
    if "draft_type" in geom_params:
        mech_rows.append(_row("Çekiş Tipi", _sanitize(geom_params.get("draft_type"))))
    if "fouling_in" in geom_params:
        mech_rows.append(_row("Boru İçi Kirlenme", f"{_fmt(geom_params.get('fouling_in'), 6)} m²K/W"))
    if "fouling_out" in geom_params:
        mech_rows.append(_row("Hava Kirlenme", f"{_fmt(geom_params.get('fouling_out'), 6)} m²K/W"))
    if "fan_diameter" in geom_params:
        mech_rows.append(_row("Fan Çapı", f"{_fmt(geom_params.get('fan_diameter'))} m"))
    if "n_fans" in geom_params:
        mech_rows.append(_row("Fan Sayısı", _sanitize(geom_params.get("n_fans"))))
    if "fan_rpm" in geom_params:
        mech_rows.append(_row("Fan Devri", f"{_fmt(geom_params.get('fan_rpm'), 0)} RPM"))
    sections.append(("Mekanik & Malzeme", mech_rows))

    if kompozisyon:
        comp_rows = []
        for name, info in kompozisyon.items():
            display = name.replace("_", " ").title()
            comp_rows.append(_row(display, f"{_sanitize(info.get('yuzde', 0))} % ({info.get('tip', 'Molar')})"))
        sections.append(("Kompozisyon", comp_rows))

    return sections


def export_excel(res, geom_params, kompozisyon, mode="Sizing"):
    wb = Workbook()
    header_font = Font(bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    label_font = Font(bold=True, size=10)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    title_font = Font(bold=True, size=14, color="2C3E50")

    sections = _build_common_sections(res, geom_params, kompozisyon, mode)

    for ws_title, rows in sections:
        safe_title = ws_title[:31]
        ws = wb.create_sheet(title=safe_title)

        ws.merge_cells("A1:B1")
        title_cell = ws["A1"]
        title_cell.value = f"Air Cooler Main — {ws_title}"
        title_cell.font = title_font
        ws.row_dimensions[1].height = 30

        for i, (label, value) in enumerate(rows, start=3):
            a = ws.cell(row=i, column=1, value=label)
            a.font = label_font
            a.border = thin_border
            b = ws.cell(row=i, column=2, value=value)
            b.border = thin_border
            b.alignment = Alignment(horizontal="right")
        ws.column_dimensions["A"].width = 35
        ws.column_dimensions["B"].width = 25

    if "T_gas_out_C" in res:
        ws_rating = wb.create_sheet(title="Rating Detay")
        ws_rating.merge_cells("A1:B1")
        ws_rating["A1"].value = "Rating Sonuçları"
        ws_rating["A1"].font = title_font
        rating_rows = [
            ("Çıkış Sıcaklığı (Gaz)", f"{_fmt(res.get('T_gas_out_C'))} °C"),
            ("Çıkış Sıcaklığı (Hava)", f"{_fmt(res.get('T_air_out_C'))} °C"),
            ("Etkinlik (ε)", f"{_fmt(res.get('effectiveness'), 4)}"),
            ("NTU", f"{_fmt(res.get('NTU'), 4)}"),
            ("Toplam Isı Geçiş Katsayısı", f"{_fmt(res.get('U_W_m2K'))} W/m²K"),
        ]
        for i, (l, v) in enumerate(rating_rows, start=3):
            ws_rating.cell(row=i, column=1, value=l).font = label_font
            ws_rating.cell(row=i, column=2, value=v)
        ws_rating.column_dimensions["A"].width = 35
        ws_rating.column_dimensions["B"].width = 25

    if "bolgeler" in res:
        ws_zones = wb.create_sheet(title="Soğutma Bölgeleri")
        ws_zones.merge_cells("A1:E1")
        ws_zones["A1"].value = "Soğutma Bölge Analizi"
        ws_zones["A1"].font = title_font
        headers = ["Bölge", "T Giriş (°C)", "T Çıkış (°C)", "Q (kW)", "Q Payı (%)"]
        for ci, h in enumerate(headers, 1):
            c = ws_zones.cell(row=3, column=ci, value=h)
            c.font = header_font
            c.fill = header_fill
            c.border = thin_border
        for ri, b in enumerate(res.get("bolgeler", []), start=4):
            ws_zones.cell(row=ri, column=1, value=b.get("bolge_adi", "")).border = thin_border
            ws_zones.cell(row=ri, column=2, value=f'{b.get("T_in_C", 0):.1f}').border = thin_border
            ws_zones.cell(row=ri, column=3, value=f'{b.get("T_out_C", 0):.1f}').border = thin_border
            ws_zones.cell(row=ri, column=4, value=f'{b.get("Q_kW", 0):.2f}').border = thin_border
            ws_zones.cell(row=ri, column=5, value=f'{b.get("Q_frac", 0) * 100:.1f}%').border = thin_border

    if "Kompozisyon" not in [s[0] for s in sections] and kompozisyon:
        ws_comp = wb.create_sheet(title="Kompozisyon")
        ws_comp.merge_cells("A1:C1")
        ws_comp["A1"].value = "Akışkan Bileşimi"
        ws_comp["A1"].font = title_font
        comp_headers = ["Bileşen", "Yüzde", "Tip"]
        for ci, h in enumerate(comp_headers, 1):
            c = ws_comp.cell(row=3, column=ci, value=h)
            c.font = header_font
            c.fill = header_fill
            c.border = thin_border
        for ri, (name, info) in enumerate(kompozisyon.items(), start=4):
            ws_comp.cell(row=ri, column=1, value=name.replace("_", " ").title()).border = thin_border
            ws_comp.cell(row=ri, column=2, value=info.get("yuzde", 0)).border = thin_border
            ws_comp.cell(row=ri, column=3, value=info.get("tip", "")).border = thin_border

    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_pdf(res, geom_params, kompozisyon, mode="Sizing"):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=20*mm, rightMargin=20*mm,
        topMargin=20*mm, bottomMargin=20*mm,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("Title2", parent=styles["Heading1"], fontSize=16, spaceAfter=6*mm, textColor=colors.HexColor("#2C3E50"))
    section_style = ParagraphStyle("Section", parent=styles["Heading2"], fontSize=12, spaceBefore=4*mm, spaceAfter=3*mm, textColor=colors.HexColor("#2980B9"))
    cell_style = ParagraphStyle("Cell", parent=styles["Normal"], fontSize=8, leading=10)
    note_style = ParagraphStyle("Note", parent=styles["Normal"], fontSize=7, textColor=colors.gray)

    elements = []
    elements.append(Paragraph(f"Air Cooler Main — {mode} Raporu", title_style))
    elements.append(Paragraph(f"Oluşturma: {datetime.now():%d.%m.%Y %H:%M}", note_style))
    elements.append(Spacer(1, 4*mm))

    sections = _build_common_sections(res, geom_params, kompozisyon, mode)

    for sec_title, rows in sections:
        elements.append(Paragraph(sec_title, section_style))
        table_data = [[Paragraph(f"<b>{label}</b>", cell_style), Paragraph(value, cell_style)] for label, value in rows]
        t = Table(table_data, colWidths=[80*mm, 80*mm])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F0F3F4")),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        elements.append(t)
        elements.append(Spacer(1, 3*mm))

    if "bolgeler" in res:
        elements.append(Paragraph("Soğutma Bölge Analizi", section_style))
        zone_data = [[Paragraph("<b>Bölge</b>", cell_style),
                      Paragraph("<b>T Giriş (°C)</b>", cell_style),
                      Paragraph("<b>T Çıkış (°C)</b>", cell_style),
                      Paragraph("<b>Q (kW)</b>", cell_style),
                      Paragraph("<b>Q Payı (%)</b>", cell_style)]]
        for b in res.get("bolgeler", []):
            zone_data.append([
                Paragraph(b.get("bolge_adi", ""), cell_style),
                Paragraph(f'{b.get("T_in_C", 0):.1f}', cell_style),
                Paragraph(f'{b.get("T_out_C", 0):.1f}', cell_style),
                Paragraph(f'{b.get("Q_kW", 0):.2f}', cell_style),
                Paragraph(f'{b.get("Q_frac", 0) * 100:.1f}%', cell_style),
            ])
        t = Table(zone_data, colWidths=[45*mm, 25*mm, 25*mm, 25*mm, 25*mm])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2C3E50")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        elements.append(t)

    elements.append(Spacer(1, 5*mm))
    elements.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#BDC3C7")))
    elements.append(Paragraph(f"Bu rapor Air Cooler Main v{APP_VERSION} tarafından otomatik oluşturulmuştur.", note_style))

    doc.build(elements)
    buf.seek(0)
    return buf


# ── Proje Yönetimi ──

def _projects_dir():
    p = Path.home() / "Air Cooler Main" / "projects"
    p.mkdir(parents=True, exist_ok=True)
    return p


PROJECT_FILENAME_RE = re.compile(r"^(.+)_(\d{8}_\d{4})\.json$")


def _safe_filename(name):
    safe = re.sub(r"[^\w\s-]", "", name).strip()
    if not safe:
        safe = "proje"
    return safe[:60]


def save_project(project_name, description, inputs, results, saved_by="user"):
    now = datetime.now()
    safe_name = _safe_filename(project_name)
    ts = now.strftime("%Y%m%d_%H%M")
    filename = f"{safe_name}_{ts}.json"
    data = {
        "version": "2.0",
        "project_name": project_name.strip(),
        "description": description.strip(),
        "saved_by": saved_by,
        "created_at": now.isoformat(),
        "updated_at": now.isoformat(),
        "type": "sizing" if results else "unknown",
        "inputs": _json_safe(inputs),
        "results": _json_safe(results or {}),
    }
    path = _projects_dir() / filename
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return str(path)


def list_projects():
    projects_dir = _projects_dir()
    if not projects_dir.exists():
        return []
    entries = []
    for fpath in sorted(projects_dir.iterdir(), key=lambda p: p.stat().st_mtime, reverse=True):
        if fpath.suffix != ".json":
            continue
        try:
            data = json.loads(fpath.read_text(encoding="utf-8"))
            entries.append({
                "filename": fpath.name,
                "project_name": data.get("project_name", fpath.stem),
                "description": data.get("description", ""),
                "saved_by": data.get("saved_by", ""),
                "created_at": data.get("created_at", ""),
                "updated_at": data.get("updated_at", ""),
                "type": data.get("type", ""),
                "path": str(fpath),
            })
        except Exception:
            entries.append({
                "filename": fpath.name,
                "project_name": fpath.stem,
                "description": "(cannot read)",
                "saved_by": "",
                "created_at": "",
                "updated_at": "",
                "type": "",
                "path": str(fpath),
            })
    return entries


def load_project_filepath(filepath):
    path = Path(filepath)
    data = json.loads(path.read_text(encoding="utf-8"))
    return data


def delete_project_file(filepath):
    path = Path(filepath)
    if path.exists():
        path.unlink()
        return True
    return False
